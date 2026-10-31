"""Import et export de l'inventaire (CSV / Excel)."""
import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from ..extensions import db, fold
from ..models import (Category, CategoryType, Condition, CustomField, Item, Location, Tracking)
from .activity import log
from .items import parse_tags

MAX_ROWS = 5000
MAX_AMOUNT = Decimal("100000000")  # colonnes Numeric(10, 2)

# champ → (libellé, synonymes reconnus dans les en-têtes)
FIELDS = {
    "code": ("Code", ["code", "code inventaire", "n inventaire", "numero inventaire", "reference", "ref"]),
    "name": ("Désignation", ["designation", "nom", "objet", "instrument", "article", "libelle"]),
    "category": ("Catégorie", ["categorie", "famille", "type", "pupitre"]),
    "brand": ("Marque", ["marque", "fabricant"]),
    "model": ("Modèle", ["modele"]),
    "serial_number": ("N° de série", ["n de serie", "numero de serie", "no de serie", "serie", "serial",
                                      "n serie", "num serie"]),
    "condition": ("État", ["etat", "condition"]),
    "location": ("Lieu", ["lieu", "emplacement", "rangement", "localisation"]),
    "tags": ("Tags", ["tags", "etiquettes", "mots cles"]),
    "purchase_date": ("Date d'achat", ["date d achat", "date achat", "achat le", "date"]),
    "purchase_price": ("Prix d'achat", ["prix d achat", "prix achat", "prix", "cout"]),
    "estimated_value": ("Valeur estimée", ["valeur estimee", "valeur", "valeur assurance", "estimation"]),
    "quantity": ("Quantité", ["quantite", "qte", "nombre", "stock"]),
    "unit": ("Unité", ["unite"]),
    "description": ("Description", ["description", "remarques", "remarque", "commentaire", "notes", "observations"]),
}
EXPORT_COLUMNS = ["code", "name", "category", "brand", "model", "serial_number", "condition", "status", "location",
                  "kit", "tags", "purchase_date", "purchase_price", "estimated_value", "quantity", "unit",
                  "description"]
EXPORT_LABELS = {**{k: v[0] for k, v in FIELDS.items()}, "status": "Statut", "kit": "Kit"}


class TransferError(ValueError):
    pass


def _normalise_header(text):
    text = fold(str(text or "")).replace("°", " ").replace("'", " ").replace("’", " ")
    return " ".join(re.sub(r"[^a-z0-9]+", " ", text).split())


# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------

def export_rows(items):
    custom_fields = db.session.scalars(db.select(CustomField).order_by(CustomField.position, CustomField.label)).all()
    header = [EXPORT_LABELS[c] for c in EXPORT_COLUMNS] + [f.label for f in custom_fields]
    rows = [header]
    for item in items:
        values = {
            "code": item.code, "name": item.name, "category": item.category.name, "brand": item.brand,
            "model": item.model, "serial_number": item.serial_number, "condition": item.condition,
            "status": item.status, "location": item.location.path if item.location else None,
            "kit": item.kit.name if item.kit else None, "tags": ", ".join(t.name for t in item.tags),
            "purchase_date": item.purchase_date, "purchase_price": item.purchase_price,
            "estimated_value": item.estimated_value, "quantity": item.quantity,
            "unit": item.unit, "description": item.description,
        }
        custom = {v.field_id: v.value for v in item.custom_values}
        rows.append([values[c] for c in EXPORT_COLUMNS] + [custom.get(f.id) for f in custom_fields])
    return rows


def to_csv(rows):
    out = io.StringIO()
    writer = csv.writer(out, delimiter=";")
    for row in rows:
        writer.writerow([_csv_cell(v) for v in row])
    return "﻿" + out.getvalue()  # BOM : Excel lit les accents


def _csv_cell(value):
    if value is None:
        return ""
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, Decimal):
        return f"{value:.2f}".replace(".", ",")
    text = str(value)
    # Neutralise les formules quand le fichier est ouvert dans un tableur.
    return "'" + text if text[:1] in ("=", "+", "-", "@", "\t", "\r") else text


def to_xlsx(rows, title="Inventaire"):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    for row in rows:
        ws.append([_xlsx_cell(v) for v in row])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for column in ws.columns:
        width = max(len(str(c.value or "")) for c in column)
        ws.column_dimensions[column[0].column_letter].width = min(max(10, width + 2), 50)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _xlsx_cell(value):
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


# ---------------------------------------------------------------------------
# Import
# ---------------------------------------------------------------------------

def read_table(data, filename):
    """Lit un CSV ou un XLSX → (en-têtes, lignes de chaînes)."""
    name = filename.lower()
    if name.endswith(".xlsx"):
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 — fichier corrompu ou pas un xlsx
            raise TransferError("Fichier Excel illisible.") from exc
        raw = [[_cell_text(c) for c in row] for row in wb.active.iter_rows(values_only=True)]
    elif name.endswith(".csv"):
        for encoding in ("utf-8-sig", "cp1252"):
            try:
                text = data.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise TransferError("Encodage du fichier non reconnu : enregistrez-le en CSV UTF-8.")
        sample = text[:4096]
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","
        try:
            raw = [row for row in csv.reader(io.StringIO(text.replace("\x00", ""), newline=""), delimiter=delimiter)]
        except csv.Error as exc:
            raise TransferError(f"CSV illisible : {exc}.") from exc
    else:
        raise TransferError("Format non pris en charge : envoyez un fichier .csv ou .xlsx.")
    raw = [[(c or "").strip() for c in row] for row in raw]
    raw = [row for row in raw if any(row)]
    if not raw:
        raise TransferError("Le fichier est vide.")
    header, rows = raw[0], raw[1:]
    if len(rows) > MAX_ROWS:
        raise TransferError(f"Trop de lignes ({len(rows)}) : maximum {MAX_ROWS} par import.")
    width = len(header)
    rows = [(row + [""] * width)[:width] for row in rows]
    return header, rows


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def guess_mapping(header):
    """Associe chaque champ à l'index de la colonne qui lui ressemble le plus (ou None)."""
    normalised = [_normalise_header(h) for h in header]
    mapping, used = {}, set()
    for field, (label, synonyms) in FIELDS.items():
        candidates = [_normalise_header(label)] + synonyms
        mapping[field] = None
        for synonym in candidates:
            for index, h in enumerate(normalised):
                if index not in used and h == synonym:
                    mapping[field] = index
                    used.add(index)
                    break
            if mapping[field] is not None:
                break
    return mapping


def _parse_date(text):
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"date invalide « {text} »")


def _parse_amount(text):
    cleaned = text.replace("€", "").replace(" ", "").replace("\xa0", "").replace(" ", "").replace(",", ".")
    try:
        value = Decimal(cleaned)
    except InvalidOperation:
        raise ValueError(f"montant invalide « {text} »") from None
    if not value.is_finite():
        raise ValueError(f"montant invalide « {text} »")
    if value < 0:
        raise ValueError(f"montant négatif « {text} »")
    if value >= MAX_AMOUNT:
        raise ValueError(f"montant trop élevé « {text} »")
    return value.quantize(Decimal("0.01"))


def _parse_int(text):
    try:
        value = int(Decimal(text.replace(",", ".")))
    except (InvalidOperation, ValueError, OverflowError):
        raise ValueError(f"quantité invalide « {text} »") from None
    if value > 1_000_000:
        raise ValueError(f"quantité trop élevée « {text} »")
    if value < 0:
        raise ValueError(f"quantité négative « {text} »")
    return value


def _condition(text):
    wanted = fold(text)
    for c in Condition.ALL:
        if fold(c) == wanted:
            return c
    synonyms = {"tres bon": "bon", "bon etat": "bon", "moyen": "usé", "use": "usé", "abime": "usé",
                "casse": "HS", "hors service": "HS", "a reparer": "à réparer", "neuve": "neuf"}
    if wanted in synonyms:
        return synonyms[wanted]
    raise ValueError(f"état inconnu « {text} » (attendu : {', '.join(Condition.ALL)})")


class _Resolver:
    """Retrouve ou crée catégories et lieux, en cache pendant un import."""

    def __init__(self, create_missing):
        self.create_missing = create_missing
        self.categories = {fold(c.name): c for c in db.session.scalars(db.select(Category))}
        self.locations = {}
        for loc in db.session.scalars(db.select(Location)):
            self.locations[fold(loc.path)] = loc

    def category(self, name):
        key = fold(name)
        if key not in self.categories:
            if not self.create_missing:
                raise ValueError(f"catégorie inconnue « {name} »")
            self.categories[key] = Category(name=name, type=CategoryType.INSTRUMENT)
            db.session.add(self.categories[key])
        return self.categories[key]

    def location(self, path):
        parts = [p.strip() for p in re.split(r"\s*>\s*|\s*/\s*", path) if p.strip()]
        parent, current = None, []
        for part in parts:
            current.append(part)
            key = fold(" > ".join(current))
            if key not in self.locations:
                if not self.create_missing:
                    raise ValueError(f"lieu inconnu « {path} »")
                self.locations[key] = Location(name=part, parent=parent)
                db.session.add(self.locations[key])
            parent = self.locations[key]
        return parent


def run_import(rows, mapping, user, *, default_category=None, create_missing=True, update_existing=True,
               prefix):
    """Importe les lignes. Renvoie (créés, mis à jour, erreurs [(n° de ligne, message)]).

    L'appelant valide (commit) ou annule (rollback, pour un essai à blanc).
    """
    resolver = _Resolver(create_missing)
    created = updated = 0
    errors = []
    seen_codes = set()

    def get(row, field):
        index = mapping.get(field)
        return row[index].strip() if index is not None and index < len(row) else ""

    for number, row in enumerate(rows, start=2):  # ligne 1 = en-têtes
        try:
            with db.session.begin_nested():
                code = get(row, "code").upper()
                if code and code in seen_codes:
                    raise ValueError(f"code {code} présent plusieurs fois dans le fichier")
                item = db.session.scalar(db.select(Item).where(db.func.upper(Item.code) == code)
                                         .execution_options(include_deleted=True)) if code else None
                if item is not None and item.deleted_at is not None:
                    raise ValueError(f"le code {code} est celui d'un objet de la corbeille (restaurez-le d'abord)")
                if item is not None and not update_existing:
                    raise ValueError(f"le code {code} existe déjà")
                creating = item is None
                name = get(row, "name")
                if creating and not name:
                    raise ValueError("désignation manquante")
                category_name = get(row, "category")
                if creating and not category_name and default_category is None:
                    raise ValueError("catégorie manquante")
                if creating:
                    item = Item(code=code or Item.next_code(prefix))
                    db.session.add(item)
                if name:
                    item.name = name[:160]
                if category_name:
                    item.category = resolver.category(category_name[:80])
                elif creating:
                    item.category = default_category
                for field in ("brand", "model", "serial_number", "unit"):
                    if value := get(row, field):
                        setattr(item, field, value[:80])
                if value := get(row, "description"):
                    item.description = value
                if value := get(row, "condition"):
                    item.condition = _condition(value)
                if value := get(row, "location"):
                    item.location = resolver.location(value)
                if value := get(row, "purchase_date"):
                    item.purchase_date = _parse_date(value)
                for field in ("purchase_price", "estimated_value"):
                    if value := get(row, field):
                        setattr(item, field, _parse_amount(value))
                if item.category.type == CategoryType.CONSOMMABLE:
                    item.tracking = Tracking.QUANTITE
                if value := get(row, "quantity"):
                    quantity = _parse_int(value)
                    if item.tracking == Tracking.QUANTITE or quantity != 1:
                        item.tracking = Tracking.QUANTITE
                        item.quantity = quantity
                if value := get(row, "tags"):
                    for tag in parse_tags(value):
                        if tag not in item.tags:
                            item.tags.append(tag)
                db.session.flush()
                seen_codes.add(item.code.upper())
                log("item.import", "Créé par import" if creating else "Mis à jour par import", item=item, user=user)
        except (ValueError, TransferError) as exc:
            errors.append((number, str(exc)))
            continue
        if creating:
            created += 1
        else:
            updated += 1
    return created, updated, errors
