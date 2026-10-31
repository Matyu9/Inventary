import csv
import io
import re
from datetime import date, datetime, timedelta
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from app.extensions import db
from app.models import (Activity, Category, CategoryType, CustomField, CustomFieldValue, Item, ItemStatus, Kit,
                        Loan, LoanLine, LoanStatus, LoanType, Location, Maintenance, Role, Tag, Tracking)
from app.services import reports, transfer

A, P, E = Role.ADMIN, Role.PROF, Role.EMPRUNTEUR
NOW = datetime.now().replace(microsecond=0)


@pytest.fixture
def park(app):
    bois, cuivres = Category(name="Bois"), Category(name="Cuivres")
    conso = Category(name="Consommables", type=CategoryType.CONSOMMABLE)
    local = Location(name="Local")
    kit = Kit(name="Kit cuivres")
    items = dict(
        sax=Item(code="OHN-0001", name="Sax", category=bois, estimated_value=Decimal("2500"), location=local),
        flute=Item(code="OHN-0002", name="Flûte", category=bois, purchase_price=Decimal("800")),
        tuba=Item(code="OHN-0003", name="Tuba", category=cuivres, estimated_value=Decimal("4000"), kit=kit),
        old=Item(code="OHN-0004", name="Vieux", category=cuivres, estimated_value=Decimal("100"),
                 status=ItemStatus.REFORME),
        cor=Item(code="OHN-0005", name="Cor", category=cuivres),
        reeds=Item(code="OHN-0006", name="Anches", category=conso, tracking=Tracking.QUANTITE, quantity=12),
    )
    db.session.add_all([*items.values(), kit, local])
    db.session.commit()
    for item in items.values():
        item.created_at = NOW - timedelta(days=400)
    db.session.commit()
    return dict(items, bois=bois, cuivres=cuivres, conso=conso, kit=kit, local=local)


def add_loan(users, status, start, end=None, items=(), kits=(), returned_at=None, loan_type=LoanType.RESERVATION):
    loan = Loan(type=loan_type, status=status, borrower=users[E], guarantor=users[P], requested_by=users[P],
                start=start, end=end, returned_at=returned_at)
    loan.lines = [LoanLine(item=i, quantity=1) for i in items] + [LoanLine(kit=k, quantity=1) for k in kits]
    db.session.add(loan)
    db.session.commit()
    return loan


# ---------------------------------------------------------------------------
# Calculs des rapports
# ---------------------------------------------------------------------------

def test_value_by_category(app, park):
    rows = {r["category"].name: r for r in reports.value_by_category()}
    assert rows["Cuivres"]["value"] == Decimal("4000") and rows["Cuivres"]["count"] == 2  # réformé exclu
    assert rows["Cuivres"]["unvalued"] == 1
    assert rows["Bois"]["value"] == Decimal("3300")  # valeur estimée ou, à défaut, prix d'achat
    assert rows["Consommables"]["count"] == 12
    assert [r["category"].name for r in reports.value_by_category()][0] == "Cuivres"


def test_usage_counts_kits_and_ignores_unshipped(app, users, park):
    add_loan(users, LoanStatus.RENDU, NOW - timedelta(days=30), NOW - timedelta(days=29), items=[park["sax"]])
    add_loan(users, LoanStatus.EN_COURS, NOW - timedelta(days=5), NOW + timedelta(days=5), items=[park["sax"]],
             kits=[park["kit"]])
    add_loan(users, LoanStatus.VALIDE, NOW + timedelta(days=1), NOW + timedelta(days=2), items=[park["flute"]])
    add_loan(users, LoanStatus.RENDU, NOW - timedelta(days=900), NOW - timedelta(days=899), items=[park["flute"]])
    rows = reports.usage(NOW - timedelta(days=365), NOW + timedelta(days=1))
    assert [(r["item"].code, r["loans"]) for r in rows] == [("OHN-0001", 2), ("OHN-0003", 1)]


def test_idle_items(app, users, park):
    add_loan(users, LoanStatus.RENDU, NOW - timedelta(days=30), NOW - timedelta(days=29), items=[park["sax"]])
    add_loan(users, LoanStatus.EN_COURS, NOW - timedelta(days=600), None, items=[park["flute"]],
             loan_type=LoanType.GARDE)
    idle = [i.code for i in reports.idle_items(NOW - timedelta(days=365))]
    assert idle == ["OHN-0003", "OHN-0005"]  # ni réformé, ni consommable, ni garde ouverte


def test_reports_use_actual_return_date(app, users, park):
    """Une garde sans date de fin, rendue il y a longtemps, ne compte plus ; un retard toujours sorti, si."""
    add_loan(users, LoanStatus.RENDU, NOW - timedelta(days=900), None, items=[park["flute"]],
             returned_at=NOW - timedelta(days=700), loan_type=LoanType.GARDE)
    add_loan(users, LoanStatus.EN_COURS, NOW - timedelta(days=500), NOW - timedelta(days=400), items=[park["sax"]])
    start, end = NOW - timedelta(days=365), NOW
    assert [r["item"].code for r in reports.usage(start, end)] == ["OHN-0001"]
    assert reports.kpis(start, end)["loans_period"] == 1
    assert [i.code for i in reports.idle_items(start)] == ["OHN-0002", "OHN-0003", "OHN-0005"]


def test_returned_late_counts_until_return(app, users, park):
    """Rendu après la fin prévue : le prêt compte jusqu'au retour effectif."""
    add_loan(users, LoanStatus.RENDU, NOW - timedelta(days=60), NOW - timedelta(days=50), items=[park["sax"]],
             returned_at=NOW - timedelta(days=20))
    assert [r["item"].code for r in reports.usage(NOW - timedelta(days=30), NOW)] == ["OHN-0001"]


def test_idle_excludes_recent_items(app, park):
    park["cor"].created_at = NOW
    db.session.commit()
    assert "OHN-0005" not in [i.code for i in reports.idle_items(NOW - timedelta(days=30))]


def test_returns_compliance(app, users, park):
    end = NOW - timedelta(days=10)
    add_loan(users, LoanStatus.RENDU, end - timedelta(days=1), end, items=[park["sax"]], returned_at=end)
    add_loan(users, LoanStatus.RENDU, end - timedelta(days=1), end, items=[park["flute"]],
             returned_at=end + timedelta(hours=20))  # dans la marge
    late = add_loan(users, LoanStatus.RENDU, end - timedelta(days=1), end, items=[park["cor"]],
                    returned_at=end + timedelta(days=3))
    result = reports.returns_compliance(NOW - timedelta(days=30), NOW)
    assert result["total"] == 3 and result["on_time"] == 2 and result["rate"] == 67 and result["late"] == [late]


def test_returns_compliance_empty(app):
    assert reports.returns_compliance(NOW - timedelta(days=30), NOW)["rate"] is None


def test_custody_grouped(app, users, park):
    other = Loan(type=LoanType.GARDE, status=LoanStatus.EN_COURS, borrower=users[P], requested_by=users[A],
                 start=NOW - timedelta(days=3))
    other.lines = [LoanLine(item=park["cor"], quantity=1)]
    db.session.add(other)
    add_loan(users, LoanStatus.EN_COURS, NOW - timedelta(days=10), None, items=[park["sax"]],
             loan_type=LoanType.GARDE)
    add_loan(users, LoanStatus.EN_COURS, NOW - timedelta(days=1), NOW + timedelta(days=1), items=[park["flute"]])
    groups = reports.custody()
    assert [(b.role, len(l)) for b, l in groups] == [(E, 1), (P, 1)] or [(b.role, len(l)) for b, l in groups] == [(P, 1), (E, 1)]
    assert all(l.type == LoanType.GARDE for _, loans in groups for l in loans)


def test_maintenance_costs(app, park):
    db.session.add_all([
        Maintenance(item=park["sax"], date=date.today() - timedelta(days=10), kind="r", cost=Decimal("120.50")),
        Maintenance(item=park["flute"], date=date.today() - timedelta(days=5), kind="r", cost=Decimal("30")),
        Maintenance(item=park["tuba"], date=date.today() - timedelta(days=5), kind="r"),
        Maintenance(item=park["tuba"], date=date.today() - timedelta(days=800), kind="r", cost=Decimal("999")),
    ])
    db.session.commit()
    rows = reports.maintenance_costs(NOW - timedelta(days=365), NOW + timedelta(days=1))
    assert rows == [{"category": "Bois", "count": 2, "cost": Decimal("150.50")},
                    {"category": "Cuivres", "count": 1, "cost": Decimal("0")}]


@pytest.mark.parametrize("value,maximum,expected", [(0, 10, 0), (5, 10, 50), (10, 10, 100), (1, 1000, 1),
                                                    (3, 0, 0), (Decimal("2.5"), Decimal("10"), 25)])
def test_bar_width(value, maximum, expected):
    assert reports.bar_width(value, maximum) == expected


def test_kpis(app, users, park):
    add_loan(users, LoanStatus.EN_COURS, NOW - timedelta(days=1), NOW + timedelta(days=1), items=[park["sax"]])
    k = reports.kpis(NOW - timedelta(days=30), NOW + timedelta(days=1))
    assert k["items"] == 5 and k["value"] == Decimal("7300") and k["active_loans"] == 1 and k["loans_period"] == 1


# ---------------------------------------------------------------------------
# Pages de rapports
# ---------------------------------------------------------------------------

def test_reports_page(client, as_role, users, park):
    add_loan(users, LoanStatus.EN_COURS, NOW - timedelta(days=5), None, items=[park["sax"]], loan_type=LoanType.GARDE)
    db.session.add(Maintenance(item=park["flute"], date=date.today(), kind="r", cost=Decimal("45")))
    db.session.commit()
    as_role(P)
    r = client.get("/rapports")
    assert r.status_code == 200
    assert "7 300,00 €" in r.text and 'class="hbar"' in r.text and "Gardes en cours par personne" in r.text
    assert "OHN-0001 Sax" in r.text and "45,00 €" in r.text
    widths = [int(w) for w in re.findall(r'class="hbar" style="width: (\d+)%"', r.text)]
    assert 100 in widths and all(0 <= w <= 100 for w in widths)


def test_reports_empty_database(client, as_role):
    as_role(A)
    r = client.get("/rapports")
    assert r.status_code == 200 and "Aucun objet." in r.text


@pytest.mark.parametrize("params", [{"debut": "2026-01-01", "fin": "2025-01-01"}, {"debut": "nope", "fin": ""},
                                    {"debut": "2026-02-30"}])
def test_reports_period_robust(client, as_role, park, params):
    as_role(P)
    assert client.get("/rapports", query_string=params).status_code == 200


def test_reports_access(client, as_role, park):
    as_role(E)
    assert client.get("/rapports").status_code == 403
    assert client.get("/rapports/valeur.csv").status_code == 403


@pytest.mark.parametrize("name,first", [("valeur", "Catégorie"), ("utilisation", "Code"), ("inactifs", "Code"),
                                        ("gardes", "Emprunteur"), ("entretiens", "Catégorie")])
def test_report_csv(client, as_role, users, park, name, first):
    add_loan(users, LoanStatus.EN_COURS, NOW - timedelta(days=5), None, items=[park["sax"]], loan_type=LoanType.GARDE)
    as_role(P)
    r = client.get(f"/rapports/{name}.csv")
    assert r.mimetype == "text/csv"
    rows = list(csv.reader(io.StringIO(r.text.lstrip("﻿")), delimiter=";"))
    assert rows[0][0] == first


def test_report_csv_unknown(client, as_role):
    as_role(P)
    assert client.get("/rapports/secret.csv").status_code == 404


# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------

def test_export_rows_with_custom_fields(app, park):
    field = CustomField(label="Tonalité", type="texte")
    db.session.add(field)
    db.session.flush()
    park["tuba"].custom_values.append(CustomFieldValue(field_id=field.id, value="Sib"))
    park["sax"].tags = [Tag(name="soliste"), Tag(name="étude")]
    db.session.commit()
    rows = transfer.export_rows(db.session.scalars(db.select(Item).order_by(Item.code)))
    header = rows[0]
    assert header[0] == "Code" and header[-1] == "Tonalité"
    tuba = rows[3]
    assert tuba[0] == "OHN-0003" and tuba[header.index("Kit")] == "Kit cuivres" and tuba[-1] == "Sib"
    assert set(rows[1][header.index("Tags")].split(", ")) == {"soliste", "étude"}
    assert rows[1][header.index("Lieu")] == "Local"


def test_csv_neutralises_formulas():
    out = transfer.to_csv([["a"], ["=HYPERLINK(\"http://x\")"], ["+1"], ["@SUM(A1)"], ["normal"],
                           [Decimal("12.5")], [date(2026, 1, 2)], [None]])
    rows = list(csv.reader(io.StringIO(out.lstrip("﻿")), delimiter=";"))
    assert [r[0] for r in rows] == ["a", "'=HYPERLINK(\"http://x\")", "'+1", "'@SUM(A1)", "normal", "12,50",
                                    "02/01/2026", ""]


def test_xlsx_roundtrip_types():
    data = transfer.to_xlsx([["Code", "Prix", "Date", "Note"], ["OHN-1", Decimal("10.5"), date(2026, 1, 2), "=1+1"]])
    ws = load_workbook(io.BytesIO(data)).active
    assert ws["B2"].value == 10.5 and ws["C2"].value.date() == date(2026, 1, 2) and ws["D2"].value == "'=1+1"
    assert ws["A1"].font.bold and ws.freeze_panes == "A2"


@pytest.mark.parametrize("fmt,mimetype", [("csv", "text/csv"),
                                          ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")])
def test_export_routes(client, as_role, park, fmt, mimetype):
    as_role(P)
    r = client.get(f"/export/inventaire.{fmt}")
    assert r.status_code == 200 and r.mimetype == mimetype and "attachment" in r.headers["Content-Disposition"]


def test_export_respects_filters(client, as_role, park):
    as_role(P)
    r = client.get("/export/inventaire.csv?category=%d" % park["bois"].id)
    rows = list(csv.reader(io.StringIO(r.text.lstrip("﻿")), delimiter=";"))
    assert [row[0] for row in rows[1:]] == ["OHN-0001", "OHN-0002"]


def test_export_access(client, as_role, park):
    as_role(E)
    assert client.get("/export/inventaire.csv").status_code == 403
    as_role(P)
    assert client.get("/export/inventaire.pdf").status_code == 404
    assert client.get("/import-export").status_code == 200
    assert "Importer depuis un tableur" not in client.get("/import-export").text


# ---------------------------------------------------------------------------
# Lecture des fichiers à importer
# ---------------------------------------------------------------------------

def xlsx_bytes(rows):
    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_read_csv_semicolon_and_bom():
    header, rows = transfer.read_table("﻿Désignation;Catégorie\nSax;Bois\n;\nFlûte;Bois\n".encode(), "x.CSV")
    assert header == ["Désignation", "Catégorie"] and rows == [["Sax", "Bois"], ["Flûte", "Bois"]]


def test_read_csv_comma_and_cp1252():
    header, rows = transfer.read_table("Nom,Famille,Prix\nFlûte,Bois,\"1 200,50\"\n".encode("cp1252"), "x.csv")
    assert header == ["Nom", "Famille", "Prix"] and rows == [["Flûte", "Bois", "1 200,50"]]


def test_read_csv_pads_short_rows():
    _, rows = transfer.read_table(b"a;b;c\n1\n1;2;3;4\n", "x.csv")
    assert rows == [["1", "", ""], ["1", "2", "3"]]


def test_read_xlsx_types():
    data = xlsx_bytes([["Nom", "Date", "Qté", "Prix"], ["Sax", datetime(2020, 5, 1), 3.0, 12.5], [None, None, None, None]])
    header, rows = transfer.read_table(data, "inv.xlsx")
    assert rows == [["Sax", "2020-05-01", "3", "12.5"]]


@pytest.mark.parametrize("data,name,message", [
    (b"", "x.csv", "vide"), (b"\n\n", "x.csv", "vide"), (b"pas un zip", "x.xlsx", "illisible"),
    (b"x", "x.ods", "Format non pris en charge"),
])
def test_read_errors(data, name, message):
    with pytest.raises(transfer.TransferError, match=message):
        transfer.read_table(data, name)


def test_read_too_many_rows(monkeypatch):
    monkeypatch.setattr(transfer, "MAX_ROWS", 2)
    with pytest.raises(transfer.TransferError, match="Trop de lignes"):
        transfer.read_table(b"a\n1\n2\n3\n", "x.csv")


def test_guess_mapping():
    header = ["N° de série", "Désignation", "Famille", "Marque", "État", "Emplacement", "Valeur assurance",
              "Remarques", "Colonne inconnue", "Qté"]
    m = transfer.guess_mapping(header)
    assert m["serial_number"] == 0 and m["name"] == 1 and m["category"] == 2 and m["brand"] == 3
    assert m["condition"] == 4 and m["location"] == 5 and m["estimated_value"] == 6 and m["description"] == 7
    assert m["quantity"] == 9 and m["code"] is None and m["tags"] is None


def test_guess_mapping_uses_each_column_once():
    m = transfer.guess_mapping(["Prix", "Prix"])
    assert m["purchase_price"] == 0 and m["estimated_value"] is None


# ---------------------------------------------------------------------------
# Import
# ---------------------------------------------------------------------------

def mapping_for(header):
    return transfer.guess_mapping(header)


def run(rows, header, users, **kw):
    kw.setdefault("prefix", "OHN")
    return transfer.run_import(rows, mapping_for(header), users[A], **kw)


def test_import_creates_items(app, users, park):
    header = ["Désignation", "Catégorie", "État", "Lieu", "Prix d'achat", "Date d'achat", "Tags", "N° de série"]
    rows = [["Clarinette", "bois", "Très bon", "Local > Armoire > Haut", "1 200,50 €", "15/09/2020", "étude, prêt", "C1"],
            ["Batterie", "Percussions", "usé", "Salle", "", "2021-01-02", "", ""]]
    created, updated, errors = run(rows, header, users)
    db.session.commit()
    assert (created, updated, errors) == (2, 0, [])
    clar = db.session.scalar(db.select(Item).filter_by(name="Clarinette"))
    assert clar.code == "OHN-0007" and clar.category == park["bois"] and clar.condition == "bon"
    assert clar.location.path == "Local > Armoire > Haut" and clar.purchase_price == Decimal("1200.50")
    assert clar.purchase_date == date(2020, 9, 15) and {t.name for t in clar.tags} == {"étude", "prêt"}
    bat = db.session.scalar(db.select(Item).filter_by(name="Batterie"))
    assert bat.category.name == "Percussions" and bat.location.name == "Salle" and bat.condition == "usé"
    assert db.session.scalar(db.select(db.func.count(Activity.id)).filter_by(action="item.import")) == 2


def test_import_reuses_existing_location_tree(app, users, park):
    run([["X", "Bois", "local / Armoire"]], ["Désignation", "Catégorie", "Lieu"], users)
    db.session.commit()
    assert db.session.scalar(db.select(db.func.count(Location.id)).filter_by(name="Local")) == 1


def test_import_updates_by_code(app, users, park):
    created, updated, errors = run([["ohn-0001", "", "", "3000"]], ["Code", "Désignation", "Catégorie", "Valeur"], users)
    db.session.commit()
    assert (created, updated, errors) == (0, 1, [])
    db.session.refresh(park["sax"])
    assert park["sax"].estimated_value == Decimal("3000") and park["sax"].name == "Sax"


def test_import_update_disabled(app, users, park):
    _, _, errors = run([["OHN-0001", "Sax"]], ["Code", "Désignation"], users, update_existing=False)
    assert errors == [(2, "le code OHN-0001 existe déjà")]


def test_import_new_code_kept(app, users, park):
    run([["ANCIEN-12", "Cornet", "Cuivres"]], ["Code", "Désignation", "Catégorie"], users)
    db.session.commit()
    assert db.session.scalar(db.select(Item).filter_by(code="ANCIEN-12")).name == "Cornet"


def test_import_duplicate_code_in_file(app, users, park):
    _, _, errors = run([["NEW-1", "A", "Bois"], ["new-1", "B", "Bois"]], ["Code", "Désignation", "Catégorie"], users)
    assert errors == [(3, "code NEW-1 présent plusieurs fois dans le fichier")]


@pytest.mark.parametrize("row,message", [
    (["", "Bois", "", "", ""], "désignation manquante"),
    (["X", "", "", "", ""], "catégorie manquante"),
    (["X", "Bois", "cassé en deux", "", ""], "état inconnu"),
    (["X", "Bois", "", "31/02/2020", ""], "date invalide"),
    (["X", "Bois", "", "", "beaucoup"], "montant invalide"),
    (["X", "Bois", "", "", "-5"], "montant négatif"),
])
def test_import_row_errors(app, users, park, row, message):
    created, _, errors = run([row], ["Désignation", "Catégorie", "État", "Date d'achat", "Prix"], users)
    assert created == 0 and len(errors) == 1 and errors[0][0] == 2 and message in errors[0][1]


def test_import_errors_do_not_block_other_rows(app, users, park):
    created, _, errors = run([["", "Bois"], ["Ok", "Bois"]], ["Désignation", "Catégorie"], users)
    assert created == 1 and [n for n, _ in errors] == [2]


def test_import_unknown_category_without_creation(app, users, park):
    _, _, errors = run([["X", "Harpes"]], ["Désignation", "Catégorie"], users, create_missing=False)
    assert "catégorie inconnue « Harpes »" in errors[0][1]


def test_import_unknown_location_without_creation(app, users, park):
    _, _, errors = run([["X", "Bois", "Grenier"]], ["Désignation", "Catégorie", "Lieu"], users, create_missing=False)
    assert "lieu inconnu" in errors[0][1]


def test_import_default_category(app, users, park):
    created, _, errors = run([["X"]], ["Désignation"], users, default_category=park["cuivres"])
    db.session.commit()
    assert created == 1 and db.session.scalar(db.select(Item).filter_by(name="X")).category == park["cuivres"]


def test_import_quantities(app, users, park):
    run([["Anches 2,5", "Consommables", "20"], ["Pupitres", "Bois", "15"], ["Sax", "Bois", "1"]],
        ["Désignation", "Catégorie", "Quantité"], users)
    db.session.commit()
    reeds = db.session.scalar(db.select(Item).filter_by(name="Anches 2,5"))
    stands = db.session.scalar(db.select(Item).filter_by(name="Pupitres"))
    sax = db.session.scalars(db.select(Item).filter_by(name="Sax")).all()[-1]
    assert reeds.is_quantity and reeds.quantity == 20
    assert stands.is_quantity and stands.quantity == 15
    assert not sax.is_quantity


def test_import_invalid_quantity(app, users, park):
    _, _, errors = run([["X", "Bois", "trois"]], ["Désignation", "Catégorie", "Quantité"], users)
    assert "quantité invalide" in errors[0][1]


# ---------------------------------------------------------------------------
# Parcours d'import (routes)
# ---------------------------------------------------------------------------

def upload(client, content, name="inventaire.csv"):
    return client.post("/import", data={"file": (io.BytesIO(content), name)}, content_type="multipart/form-data")


def test_import_flow(client, as_role, park):
    as_role(A)
    r = upload(client, "Désignation;Catégorie;État\nClarinette;Bois;bon\nBasson;Bois;usé\n".encode())
    token = r.headers["Location"].rsplit("/", 1)[1]
    page = client.get(f"/import/{token}")
    assert "2 ligne(s) détectée(s)" in page.text and "Clarinette" in page.text
    assert re.search(r'name="map_name">.*?<option value="0" selected', page.text, re.S)
    form = {"map_name": "0", "map_category": "1", "map_condition": "2", "create_missing": "1", "update_existing": "1"}
    r = client.post(f"/import/{token}", data={**form, "dry_run": "1"})
    assert "Essai réussi : 2 objet(s)" in r.text
    assert db.session.scalar(db.select(db.func.count(Item.id))) == 6  # essai à blanc : rien d'écrit
    r = client.post(f"/import/{token}", data={**form, "confirm": "1"})
    assert r.headers["Location"].endswith("/objets")
    assert db.session.scalar(db.select(db.func.count(Item.id))) == 8
    assert client.get(f"/import/{token}").status_code == 404  # fichier temporaire supprimé


def test_import_confirm_refused_with_errors(client, as_role, park):
    as_role(A)
    token = upload(client, b"Nom;Cat\n;Bois\nOk;Bois\n").headers["Location"].rsplit("/", 1)[1]
    r = client.post(f"/import/{token}", data={"map_name": "0", "map_category": "1", "confirm": "1"})
    assert "1 ligne(s) en erreur" in r.text and "Ligne 2 : désignation manquante" in r.text
    assert db.session.scalar(db.select(db.func.count(Item.id))) == 6


def test_import_requires_name_mapping(client, as_role, park):
    as_role(A)
    token = upload(client, b"a;b\n1;2\n").headers["Location"].rsplit("/", 1)[1]
    r = client.post(f"/import/{token}", data={"dry_run": "1"})
    assert "Associez au moins" in r.text


def test_import_ignores_out_of_range_mapping(client, as_role, park):
    as_role(A)
    token = upload(client, b"Nom\nX\n").headers["Location"].rsplit("/", 1)[1]
    r = client.post(f"/import/{token}", data={"map_name": "0", "map_category": "99", "dry_run": "1"})
    assert "catégorie manquante" in r.text


def test_import_xlsx_upload(client, as_role, park):
    as_role(A)
    r = upload(client, xlsx_bytes([["Désignation", "Catégorie"], ["Hautbois", "Bois"]]), "f.xlsx")
    assert "/import/" in r.headers["Location"]


def test_import_bad_file(client, as_role):
    as_role(A)
    r = client.post("/import", data={"file": (io.BytesIO(b"x"), "a.pdf")}, content_type="multipart/form-data",
                    follow_redirects=True)
    assert "Format non pris en charge" in r.text
    r = client.post("/import", data={}, content_type="multipart/form-data", follow_redirects=True)
    assert "Choisissez un fichier" in r.text


def test_import_token_bound_to_session(client, app, as_role, users, park):
    as_role(A)
    token = upload(client, b"Nom\nX\n").headers["Location"].rsplit("/", 1)[1]
    other = app.test_client()
    other.post("/connexion", data={"email": users[A].email, "password": "motdepasse"})
    assert other.get(f"/import/{token}").status_code == 404
    assert client.get("/import/..%2f..%2fetc").status_code == 404


def test_import_admin_only(client, as_role, park):
    as_role(P)
    assert upload(client, b"Nom\nX\n").status_code == 403


def test_export_then_reimport_roundtrip(client, as_role, park):
    as_role(A)
    csv_text = client.get("/export/inventaire.csv").data
    token = upload(client, csv_text).headers["Location"].rsplit("/", 1)[1]
    page = client.get(f"/import/{token}").text
    selected = dict(re.findall(r'name="map_(\w+)">\s*<option value="">— ignorer —</option>.*?<option value="(\d+)" selected',
                               page, re.S))
    form = {f"map_{k}": v for k, v in selected.items()}
    r = client.post(f"/import/{token}", data={**form, "update_existing": "1", "dry_run": "1"})
    assert "0 objet(s) seront créés et 6 mis à jour" in r.text


def test_import_uses_given_prefix(app, users, park):
    created, _, errors = run([["Hélicon", "Cuivres"]], ["Désignation", "Catégorie"], users, prefix="ABC")
    assert (created, errors) == (1, [])
    assert db.session.scalar(db.select(Item).filter_by(name="Hélicon")).code == "ABC-0001"


def test_report_tables_scroll_on_small_screens(client, as_role, park):
    """Sur téléphone, les tableaux des rapports défilent au lieu d'être coupés."""
    as_role(A)
    page = client.get("/rapports").text
    assert '<table' in page
    assert page.count('<table') == page.count('<div class="overflow-auto"><table')
